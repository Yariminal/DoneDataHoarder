"""DOCX body extraction keeps table evidence in document order and within caps."""
import zipfile
import json

from docx import Document
from sqlalchemy.orm import Session

from donedatahoarder.analyzers import document
from donedatahoarder.analyzers.base import (
    AnalysisResult, EXTRACTOR_VERSION, PROMPT_VERSION,
)
from donedatahoarder.analyzers.cache import context_hash, extractor_key, restore
from donedatahoarder.analyzers.image import ImageAnalyzer
from donedatahoarder.db.models import AnalysisCache, File, UserSession
from donedatahoarder.db.session import get_engine, init_db


def test_mixed_docx_prompt_contains_body_table_and_following_paragraph(tmp_path):
    path = tmp_path / "schedule.docx"
    doc = Document()
    doc.add_paragraph("Coordinator: Rowan")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Session"
    table.cell(0, 1).text = "Date"
    table.cell(0, 2).text = "Contact"
    table.cell(1, 0).text = "Field survey briefing"
    table.cell(1, 1).text = "2027-04-13"
    table.cell(1, 2).text = "Avery, operations lead"
    doc.add_paragraph("Confirm travel and equipment after the briefing.")
    doc.save(path)

    prompts = []

    class Client:
        def generate_json(self, prompt, **_kwargs):
            prompts.append(prompt)
            return {"description": "Field survey schedule and contacts",
                    "suggested_name": "field_survey_schedule", "tags": ["survey"],
                    "document_type": "report", "confidence": 0.8}

    file = File(path=str(path), filename=path.name, extension=".docx")
    result = document.DocumentAnalyzer(Client()).analyze(file, "Inbox")
    assert result.content_available is True
    assert result.evidence_source == "text"
    text = document.extract_document(path).text
    assert text.index("Coordinator: Rowan") < text.index("Session | Date | Contact")
    assert text.index("Field survey briefing") < text.index("Confirm travel")
    assert "Field survey briefing | 2027-04-13 | Avery, operations lead" in prompts[0]


def test_table_only_and_nested_table_are_read_once(tmp_path):
    path = tmp_path / "contacts.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Workstream"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Inspection schedule"
    owner = table.cell(1, 1)
    owner.text = "Regional team"
    nested = owner.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Backup"
    nested.cell(0, 1).text = "Morgan"
    doc.save(path)
    extraction = document.extract_document(path)
    assert extraction.reason is None
    assert extraction.text.count("Inspection schedule") == 1
    assert extraction.text.count("Backup | Morgan") == 1


def test_merged_cells_do_not_repeat_text(tmp_path):
    path = tmp_path / "merged.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Shared heading"
    table.cell(1, 0).text = "Workshop"
    table.cell(1, 1).text = "2027-05-02"
    doc.save(path)
    assert document.extract_document(path).text == "Shared heading\nWorkshop | 2027-05-02"


def test_oversized_expanded_xml_and_malformed_docx_fail_without_model(tmp_path, monkeypatch):
    path = tmp_path / "large.docx"
    doc = Document()
    doc.add_paragraph("A substantive introduction.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Detailed planning notes in a table."
    doc.save(path)
    with zipfile.ZipFile(path) as archive:
        xml_size = archive.getinfo("word/document.xml").file_size
    monkeypatch.setattr(document, "MAX_DOCX_DOCUMENT_XML_BYTES", xml_size - 1)

    class NoModel:
        def generate_json(self, *_args, **_kwargs):
            raise AssertionError("No model call for rejected DOCX")

    file = File(path=str(path), filename=path.name, extension=".docx")
    result = document.DocumentAnalyzer(NoModel()).analyze(file, "context")
    assert result.outcome == "skipped"
    assert result.reason == "oversized_content"
    assert result.evidence_source == "none"
    monkeypatch.setattr(document, "MAX_DOCX_DOCUMENT_XML_BYTES", xml_size + 1)
    path.write_bytes(b"not a DOCX")
    malformed = document.DocumentAnalyzer(NoModel()).analyze(file, "context")
    assert malformed.outcome == "skipped"
    assert malformed.reason == "unreadable_content"
    assert malformed.evidence_source == "none"


def test_extractor_version_invalidates_prior_paragraph_only_cache():
    assert EXTRACTOR_VERSION == "extractors-v4-2026-09-28"
    assert document.DOCX_EXTRACTOR_VERSION == "extractors-v5-2026-09-28"


def test_old_docx_cache_misses_while_unrelated_image_cache_hits(tmp_path):
    init_db(tmp_path / "index.db")
    with Session(get_engine()) as db:
        user = UserSession(root_path=str(tmp_path), name="versions")
        db.add(user)
        db.flush()
        docx = File(session_id=user.id, path=str(tmp_path / "same.docx"),
                    filename="same.docx")
        image = File(session_id=user.id, path=str(tmp_path / "same.png"),
                     filename="same.png")
        db.add_all([docx, image])
        db.flush()
        text_analyzer = document.DocumentAnalyzer(None)
        image_analyzer = ImageAnalyzer(None)
        assert extractor_key(text_analyzer, docx) == (
            "DocumentAnalyzer/extractors-v5-2026-09-28")
        assert extractor_key(image_analyzer, image) == (
            "ImageAnalyzer/extractors-v4-2026-09-28")
        same_hash = "a" * 64
        same_context = context_hash("same context")
        for name, source in (("DocumentAnalyzer", "text"),
                             ("ImageAnalyzer", "vision")):
            db.add(AnalysisCache(
                content_sha256=same_hash, context_sha256=same_context,
                model_tag="local:test", model_digest="digest",
                prompt_version=PROMPT_VERSION,
                extractor_version=f"{name}/extractors-v4-2026-09-28",
                payload_json=json.dumps({"analysis_outcome": "content_verified",
                                         "analysis_evidence_source": source,
                                         "ai_description": "prior evidence"}),
            ))
        db.commit()
        assert not restore(db, docx, content_hash=same_hash,
                           context_digest=same_context,
                           model_keys=[("local:test", "digest", "text")],
                           analyzer=text_analyzer)
        assert restore(db, image, content_hash=same_hash,
                       context_digest=same_context,
                       model_keys=[("local:test", "digest", "vision")],
                       analyzer=image_analyzer)
        assert image.analysis_cache_hit
        text_analyzer.save_result(
            docx, AnalysisResult(description="table evidence", evidence_source="text",
                                 extractor="python-docx"), "local:test", "digest",
        )
        db.refresh(docx)
        assert docx.analysis_extractor_version == "python-docx/extractors-v5-2026-09-28"
